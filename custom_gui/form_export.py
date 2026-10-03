import csv
import re
import os
import argparse
import typing
import openpyxl

def spread_rows(n_pages: int) -> list[tuple[str, str, list[int]]]:
    """[(A, B, pdf_pages), ...] as in section 3."""
    if n_pages < 6 or n_pages % 2 != 0:
        raise ValueError("n_pages must be an even integer >= 6")
        
    rows = []
    rows.append(("01", "表紙", [1]))
    rows.append(("02", "表見返し-P.0001", [2, 3]))
    
    for k in range(3, n_pages // 2):
        rows.append((f"{k:02d}", f"P.{2*k-4:04d}-P.{2*k-3:04d}", [2*k-2, 2*k-1]))
        
    k = n_pages // 2
    rows.append((f"{k:02d}", f"P.{n_pages-4:04d}-裏見返し", [n_pages-2, n_pages-1]))
    
    k = n_pages // 2 + 1
    rows.append((f"{k:02d}", "裏表紙", [n_pages]))
    
    return rows

def read_viewer_csv(path: str) -> list[dict]:
    """Rows of the viewer CSV as dicts (utf-8-sig)."""
    with open(path, mode="r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        return list(reader)

def build_keywords(csv_rows: list[dict], n_pages: int) -> list[str | None]:
    """Column C for every spread, same length as spread_rows(n_pages)."""
    page_rows = {}
    stem_set = set()
    
    pattern = re.compile(r"_p?(\d+)\.[A-Za-z0-9]+$")
    
    for row in csv_rows:
        image_name = row["image_name"]
        match = pattern.search(image_name)
        if not match:
            raise ValueError(f"Invalid image_name: {image_name}")
        page_num = int(match.group(1))
        stem = image_name[:match.start()]
        
        stem_set.add(stem)
        if len(stem_set) > 1:
            raise ValueError("Multiple stems found in CSV")
            
        if page_num < 1 or page_num > n_pages:
            raise ValueError(f"Page number out of bounds: {page_num}")
            
        if page_num not in page_rows:
            page_rows[page_num] = []
        page_rows[page_num].append(row)
        
    s_rows = spread_rows(n_pages)
    keywords = []
    
    tag_pattern = re.compile(r"^【[^【】]+】$")
    
    for spread_idx, (_, _, pdf_pages) in enumerate(s_rows):
        spread_segments = []
        
        for page in sorted(pdf_pages):
            rows_for_page = page_rows.get(page, [])
            
            mark_texts = []
            region_texts = []
            
            for r in rows_for_page:
                if r["region_id"] == "mark":
                    mark_texts.append(r["text"])
                elif r["region_id"] != "1":
                    region_texts.append(r["text"])
            
            for t in mark_texts + region_texts:
                if not t:
                    continue
                t = t.replace("\r\n", "｜").replace("\n", "｜")
                for piece in t.split("｜"):
                    piece = piece.strip()
                    if piece:
                        spread_segments.append(piece)
                        
        final_segments = []
        seen_tags = set()
        
        for seg in spread_segments:
            if tag_pattern.match(seg):
                if seg not in seen_tags:
                    seen_tags.add(seg)
                    final_segments.append(seg)
            else:
                final_segments.append(seg)
                
        if spread_idx == 0:
            if "【表紙】" not in final_segments:
                final_segments.insert(0, "【表紙】")
                
        if final_segments:
            keywords.append("｜".join(final_segments))
        else:
            keywords.append(None)
            
    return keywords

def write_form(template_path: str, out_path: str, n_pages: int, keywords: list[str | None]) -> None:
    """Load template, fill rows from 3, save to out_path."""
    if os.path.abspath(template_path) == os.path.abspath(out_path):
        raise ValueError("template_path and out_path are the same file")
    if os.path.exists(out_path):
        raise FileExistsError(f"Output path exists: {out_path}")
        
    s_rows = spread_rows(n_pages)
    
    wb = openpyxl.load_workbook(template_path)
    ws = wb["入力フォーム"]
    
    target_max_row = max(ws.max_row, 2 + len(s_rows))
    
    for row_idx in range(3, target_max_row + 1):
        if row_idx - 3 < len(s_rows):
            A_val, B_val, _ = s_rows[row_idx - 3]
            C_val = keywords[row_idx - 3]
            ws.cell(row=row_idx, column=1).value = A_val
            ws.cell(row=row_idx, column=2).value = B_val
            ws.cell(row=row_idx, column=3).value = C_val
            ws.cell(row=row_idx, column=4).value = f'=IF(LEN(B{row_idx})+LEN(C{row_idx})<>0,LEN(B{row_idx})+LEN(C{row_idx}),"")'
        else:
            ws.cell(row=row_idx, column=1).value = None
            ws.cell(row=row_idx, column=2).value = None
            ws.cell(row=row_idx, column=3).value = None
            
    wb.save(out_path)

def main(argv=None) -> int:
    """python -m custom_gui.form_export --csv C --pages N --template T --out O"""
    parser = argparse.ArgumentParser()
    parser.add_argument("--csv", required=True)
    parser.add_argument("--pages", type=int, required=True)
    parser.add_argument("--template", required=True)
    parser.add_argument("--out", required=True)
    args = parser.parse_args(argv)
    
    try:
        csv_rows = read_viewer_csv(args.csv)
        keywords = build_keywords(csv_rows, args.pages)
        write_form(args.template, args.out, args.pages, keywords)
    except Exception as e:
        print(f"Error: {e}")
        return 1
    return 0

if __name__ == '__main__':
    import sys
    sys.exit(main())
