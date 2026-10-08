# -*- coding: utf-8 -*-
"""141号 納品 Excel の字を、承認者のご指示（10-08 22:4x）どおり位置を決めて置き換え、別名で書き出す。
元のファイルは変えない。置く前に、その位置の字が「いま」の字と一致するかを確かめ、1 つでも違えば何も書かずに止まる。
書いたあとで読み戻し、全シート・全セルを元と比べて、違いが決めた字だけであることを確かめる。"""
import os, sys
sys.stdout.reconfigure(encoding="utf-8")
import openpyxl
D = r"C:\Users\user\ndlocr-lite-custom\work\納品\261008"
SRC = os.path.join(D, "（ご納品）国際寫眞新聞_141号_p01-52_20261008.xlsx")
DST = os.path.join(D, "（ご納品／修正）国際寫眞新聞_141号_p01-52_20261008.xlsx")
SHEET = "入力フォーム"
# (セル, 位置 1 始まり, いま, 直す字)
FIXES = [
    ("C24", 839, "郷", "鄕"),  # いゝ男の西鄕さん（見出し）
    ("C24", 864, "郷", "鄕"),  # 正裝の西鄕さんの
    ("C24", 886, "児", "兒"),  # 更に鹿兒島縣出身
    ("C24", 931, "頼", "賴"),  # 製作を依賴
    ("C24", 973, "装", "裝"),  # 陸軍大將正裝、山城信國
]
if os.path.exists(DST):
    sys.exit(f"止めました: 書き出し先が既にあります: {DST}")
wb = openpyxl.load_workbook(SRC)
ws = wb[SHEET]
vals = {}
for cell, pos, now, new in FIXES:
    v = vals.get(cell, ws[cell].value)
    if v[pos - 1] != now:
        sys.exit(f"止めました: {cell} {pos} 字目は「{v[pos - 1]}」で、「{now}」ではありません（何も書いていません）")
    vals[cell] = v[:pos - 1] + new + v[pos:]
for cell, v in vals.items():
    ws[cell].value = v
wb.save(DST)

a, b = openpyxl.load_workbook(SRC), openpyxl.load_workbook(DST)
assert a.sheetnames == b.sheetnames, (a.sheetnames, b.sheetnames)
diffs = []
for name in a.sheetnames:
    sa, sb = a[name], b[name]
    assert (sa.max_row, sa.max_column) == (sb.max_row, sb.max_column), name
    for row in sa.iter_rows():
        for c in row:
            x, y = c.value, sb[c.coordinate].value
            if x == y: continue
            if not (isinstance(x, str) and isinstance(y, str) and len(x) == len(y)):
                diffs.append((name, c.coordinate, "長さか型が違う")); continue
            for i, (p, q) in enumerate(zip(x, y), 1):
                if p != q: diffs.append((name, c.coordinate, i, p, q))
want = sorted((SHEET, c, p, n, w) for c, p, n, w in FIXES)
print("違い:", len(diffs), "件")
for d in diffs: print("  ", d)
print("決めた字と一致:", sorted(diffs) == want)
