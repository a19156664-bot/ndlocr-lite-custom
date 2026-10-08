# -*- coding: utf-8 -*-
"""141号 全 52 頁の Excel を、承認者ご指示の割り振り（10-08）で作る。
PDF 1 → C3 表紙／C4 表見返し-P.0001 は空／PDF n（2..49）→ P.n の行／
末尾は案A: C28「P.0048-P.0049」← 48・49、C29「P.0050-裏見返し」← 50・51、C30「裏表紙」← 52。
リポジトリのコードは変えない。form_export の spread_rows をこのスクリプトの中だけで差し替える。"""
import os, sys, json
sys.stdout.reconfigure(encoding="utf-8")
BASE = r"C:\Users\user\ndlocr-lite-custom"
sys.path.insert(0, BASE)
OUT = os.path.dirname(os.path.abspath(__file__))
import openpyxl, pypdfium2
import custom_gui.form_export as fe
import custom_gui.work_state as work_state
from custom_gui.selection import SelectionContainer, SelectionRect
from custom_gui.exporter import build_export_rows_multi, rows_to_csv_text

ISSUE = "国際寫眞新聞_141号"
PDF = os.path.join(BASE, "work", f"{ISSUE}.pdf")
RAW = os.path.join(BASE, "work", "output", "01_raw_ocr", f"{ISSUE}.json")
TEMPLATE = os.path.join(BASE, "work", "form0823.xlsx")
NEW_XLSX = os.path.join(BASE, "work", sys.argv[1])

def rows_141(n_pages):
    assert n_pages == 52
    rows = [("01", "表紙", [1]), ("02", "表見返し-P.0001", [])]
    k = 3
    for p in range(2, 50, 2):                      # P.0002-P.0003 … P.0048-P.0049
        rows.append((f"{k:02d}", f"P.{p:04d}-P.{p+1:04d}", [p, p + 1])); k += 1
    rows.append((f"{k:02d}", "P.0050-裏見返し", [50, 51])); k += 1
    rows.append((f"{k:02d}", "裏表紙", [52]))
    return rows

def normalize(page_lines, name):
    res = []
    for b in page_lines:
        if "boundingBox" not in b or "text" not in b:
            continue
        xs, ys = [], []
        for pt in b["boundingBox"]:
            if isinstance(pt, str):
                x, y = map(float, pt.split()); xs.append(x); ys.append(y)
            elif isinstance(pt, (list, tuple)):
                xs.append(float(pt[0])); ys.append(float(pt[1]))
        if not xs:
            continue
        x1, x2, y1, y2 = min(xs), max(xs), min(ys), max(ys)
        if x1 >= x2: x2 = x1 + 1
        if y1 >= y2: y2 = y1 + 1
        v = b.get("isVertical", False)
        res.append({"text": str(b.get("text", "")), "bbox": (x1, y1, x2, y2),
                    "confidence": float(b.get("confidence", 0.0)),
                    "is_vertical": v.lower() == "true" if isinstance(v, str) else bool(v),
                    "source_image": name})
    return res

doc = pypdfium2.PdfDocument(PDF); n = len(doc); doc.close()
contents = json.load(open(RAW, encoding="utf-8")).get("contents", [])
rows, saved = [], []
for i in range(n):
    name = f"{ISSUE}_p{i+1:04d}.png"
    parsed = normalize(contents[i], name) if i < len(contents) else []
    st = work_state.load_work_state(PDF, page_index=i)
    c = SelectionContainer(); edits = {}; mark = None
    if st:
        saved.append(i + 1)
        c.restore([SelectionRect(rect_id=r["rect_id"], bbox=tuple(r["bbox"]), label=r["label"]) for r in st["rects"]])
        edits = st["edits"]; mark = st["mark"]
    if len(c.get_all()) == 0 and parsed:
        xs = [l["bbox"][0] for l in parsed] + [l["bbox"][2] for l in parsed]
        ys = [l["bbox"][1] for l in parsed] + [l["bbox"][3] for l in parsed]
        c.add((max(0.0, min(xs) - 15), max(0.0, min(ys) - 15), max(xs) + 15, max(ys) + 15))
    rects = c.get_all()
    if mark:
        rows.append({"image_name": name, "region_id": "mark", "x1": 0, "y1": 0, "x2": 0, "y2": 0, "line_count": 0, "text": mark})
    if rects:
        rows.extend(build_export_rows_multi([{"image_name": name, "rects": rects, "ocr_results": parsed, "edited_texts": edits}]))
r1_edited = [i + 1 for i in range(n) if (work_state.load_work_state(PDF, page_index=i) or {}).get("edits", {}).get("1")]
print(f"Region1 に人が文字を入れた頁（落ちる）: {r1_edited}")
print(f"pdf pages={n} saved_state pages={len(saved)} missing={[p for p in range(1, n + 1) if p not in saved]}")
csv_path = os.path.join(OUT, "export141_p1-52_final.csv")
with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
    f.write(rows_to_csv_text(rows))
print(f"csv rows={len(rows)} (region_id=1: {sum(1 for r in rows if str(r['region_id']) == '1')})")

# どのページにも納品に入る文字があるか（region_id=1 を除く）
pages_with_text = sorted({int(r["image_name"][-8:-4]) for r in rows if str(r["region_id"]) != "1" and str(r["text"]).strip()})
print(f"pages with delivered text={len(pages_with_text)} none={[p for p in range(1, n + 1) if p not in pages_with_text]}")

orig = fe.spread_rows
fe.spread_rows = rows_141
try:
    kw = fe.build_keywords(fe.read_viewer_csv(csv_path), n)
    fe.write_form(TEMPLATE, NEW_XLSX, n, kw)
finally:
    fe.spread_rows = orig
ws = openpyxl.load_workbook(NEW_XLSX)["入力フォーム"]
print(f"[出力] {NEW_XLSX}")
for r in range(3, 32):
    a, b, cc = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value
    if a is None and b is None and cc is None:
        continue
    print(f"  C{r}: A={a} B={b} C={(str(cc)[:36] if cc else '')}")
