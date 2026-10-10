# -*- coding: utf-8 -*-
"""号の Excel を、141号と同じ割り振りの決まり（承認者 10-08・make141_final.py）で作る（10-10・142号から）。
割り振り（頁数 n は偶数）: PDF 1 → 表紙／表見返し-P.0001 は空／PDF p,p+1 → P.p-P.p+1（p=2..n-4）／
PDF n-2,n-1 → P.(n-2)-裏見返し／PDF n → 裏表紙。n=52 で make141_final.py の rows_141 と同じになる。
リポジトリのコードは変えない。form_export の spread_rows をこのスクリプトの中だけで差し替える。

使い方: .venv\\Scripts\\python.exe .nightly\\verify\\make_issue_final.py <号> <出力の xlsx の名前（work の下）>
"""
import os, sys, json
sys.stdout.reconfigure(encoding="utf-8")
BASE = r"C:\Users\user\ndlocr-lite-custom"
sys.path.insert(0, BASE)
OUT = os.path.join(BASE, "work", "proofread")  # 中間の CSV はお客様の文字の全文なので、git の外（work）に置く
import openpyxl, pypdfium2
import custom_gui.form_export as fe
import custom_gui.work_state as work_state
from custom_gui.selection import SelectionContainer, SelectionRect
from custom_gui.exporter import build_export_rows_multi, rows_to_csv_text


def rows_issue(n_pages):
    assert n_pages % 2 == 0 and n_pages >= 8
    rows = [("01", "表紙", [1]), ("02", "表見返し-P.0001", [])]
    k = 3
    for p in range(2, n_pages - 2, 2):
        rows.append((f"{k:02d}", f"P.{p:04d}-P.{p+1:04d}", [p, p + 1])); k += 1
    rows.append((f"{k:02d}", f"P.{n_pages - 2:04d}-裏見返し", [n_pages - 2, n_pages - 1])); k += 1
    rows.append((f"{k:02d}", "裏表紙", [n_pages]))
    return rows


def normalize(page_lines, name):
    res = []
    for b in page_lines:
        if "boundingBox" not in b or "text" not in b:
            continue
        xs, ys = [], []
        for pt in b["boundingBox"]:
            if isinstance(pt, str):
                x, y = map(float, pt.split()); xs.append(x); ys.append(y)
            elif isinstance(pt, (list, tuple)):
                xs.append(float(pt[0])); ys.append(float(pt[1]))
        if not xs:
            continue
        x1, x2, y1, y2 = min(xs), max(xs), min(ys), max(ys)
        if x1 >= x2: x2 = x1 + 1
        if y1 >= y2: y2 = y1 + 1
        v = b.get("isVertical", False)
        res.append({"text": str(b.get("text", "")), "bbox": (x1, y1, x2, y2),
                    "confidence": float(b.get("confidence", 0.0)),
                    "is_vertical": v.lower() == "true" if isinstance(v, str) else bool(v),
                    "source_image": name})
    return res


def main():
    num, out_name = sys.argv[1], sys.argv[2]
    issue = f"国際寫眞新聞_{num}号"
    pdf = os.path.join(BASE, "work", f"{issue}.pdf")
    raw = os.path.join(BASE, "work", "output", "01_raw_ocr", f"{issue}.json")
    template = os.path.join(BASE, "work", "form0823.xlsx")
    new_xlsx = os.path.join(BASE, "work", out_name)
    doc = pypdfium2.PdfDocument(pdf); n = len(doc); doc.close()
    contents = json.load(open(raw, encoding="utf-8")).get("contents", [])
    rows, saved = [], []
    for i in range(n):
        name = f"{issue}_p{i+1:04d}.png"
        parsed = normalize(contents[i], name) if i < len(contents) else []
        st = work_state.load_work_state(pdf, page_index=i)
        c = SelectionContainer(); edits = {}; mark = None
        if st:
            saved.append(i + 1)
            c.restore([SelectionRect(rect_id=r["rect_id"], bbox=tuple(r["bbox"]), label=r["label"]) for r in st["rects"]])
            edits = st["edits"]; mark = st["mark"]
        if len(c.get_all()) == 0 and parsed:
            xs = [l["bbox"][0] for l in parsed] + [l["bbox"][2] for l in parsed]
            ys = [l["bbox"][1] for l in parsed] + [l["bbox"][3] for l in parsed]
            c.add((max(0.0, min(xs) - 15), max(0.0, min(ys) - 15), max(xs) + 15, max(ys) + 15))
        rects = c.get_all()
        if mark:
            rows.append({"image_name": name, "region_id": "mark", "x1": 0, "y1": 0, "x2": 0, "y2": 0, "line_count": 0, "text": mark})
        if rects:
            rows.extend(build_export_rows_multi([{"image_name": name, "rects": rects, "ocr_results": parsed, "edited_texts": edits}]))
    r1_edited = [i + 1 for i in range(n) if (work_state.load_work_state(pdf, page_index=i) or {}).get("edits", {}).get("1")]
    print(f"Region1 に人が文字を入れた頁（落ちる）: {r1_edited}")
    print(f"pdf pages={n} saved_state pages={len(saved)} missing={[p for p in range(1, n + 1) if p not in saved]}")
    csv_path = os.path.join(OUT, f"export{num}_p1-{n}_final.csv")
    with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
        f.write(rows_to_csv_text(rows))
    print(f"csv rows={len(rows)} (region_id=1: {sum(1 for r in rows if str(r['region_id']) == '1')})")
    pages_with_text = sorted({int(r["image_name"][-8:-4]) for r in rows if str(r["region_id"]) != "1" and str(r["text"]).strip()})
    print(f"pages with delivered text={len(pages_with_text)} none={[p for p in range(1, n + 1) if p not in pages_with_text]}")
    orig = fe.spread_rows
    fe.spread_rows = rows_issue
    try:
        kw = fe.build_keywords(fe.read_viewer_csv(csv_path), n)
        fe.write_form(template, new_xlsx, n, kw)
    finally:
        fe.spread_rows = orig
    ws = openpyxl.load_workbook(new_xlsx)["入力フォーム"]
    print(f"[出力] {new_xlsx}")
    for r in range(3, 3 + len(rows_issue(n)) + 2):
        a, b, cc = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value
        if a is None and b is None and cc is None:
            continue
        print(f"  C{r}: A={a} B={b} 字数={len(str(cc)) if cc else 0} 頭={(str(cc)[:30] if cc else '')}")


if __name__ == "__main__":
    main()
